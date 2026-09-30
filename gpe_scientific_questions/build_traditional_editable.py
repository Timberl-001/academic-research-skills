#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Draw one fully editable, original conventional-GPE problem slide.

Every visible object is a native PowerPoint shape, connector, freeform or text
box. The two schematics are not PNG pictures embedded in a slide. SVG and PPT
use the same named scene. No specific proposed electrolyte recipe appears.

Usage:
    python3 gpe_scientific_questions/build_traditional_editable.py --font-dir /tmp/fonts
"""
from __future__ import annotations

import argparse
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

import resvg_py
from pptx.enum.shapes import MSO_SHAPE_TYPE

from build_design_concept import Scene

HERE = Path(__file__).resolve().parent
BASE = HERE / "Traditional_GPE_Common_Problems_Editable"
TITLE = "关键科学问题：传统凝胶电解质的两大共性挑战"
DESCRIPTION = (
    "Original schematic of common conventional gel-polymer-electrolyte challenges: "
    "bulk ion transport and electrode/interphase stability. Qualitative examples, "
    "not universal failure mechanisms or measured structures. All elements are editable."
)
C = {
    "bg": "#F5F7FB", "white": "#FFFFFF", "ink": "#172B46",
    "muted": "#64768A", "edge": "#DCE4ED", "mesh": "#5F6E7D",
    "blue": "#2B6FA9", "blue_dark": "#225A8C", "blue_light": "#EDF5FB",
    "coral": "#BE5D50", "coral_light": "#FCF0EC",
    "solvent": "#C4E6F5", "solvent_edge": "#8BC6DF",
    "gold": "#CC9232", "node": "#D4864D",
}


def solvent(s, x, y, prefix, size=12):
    s.ellipse(x, y, size, size * 0.72, C["solvent"], C["solvent_edge"], 1,
              name=prefix + "-溶剂符号")
    s.text(x, y + 4, "S", 13, C["blue"], True, "middle",
           name=prefix + "-溶剂文字")


def node(s, x, y, prefix, radius=7):
    s.ellipse(x, y, radius, radius, C["node"], "#B36F41", 1.2,
              name=prefix + "-交联节点")
    s.ellipse(x - 2, y - 2, radius * .35, radius * .27, "#F3CDAE",
              name=prefix + "-节点高光")


def panel(s, x, number, heading, question, accent, tint, prefix):
    s.rect(x, 192, 888, 724, C["white"], C["edge"], 1.8, 22,
           name=prefix + "-面板外框")
    s.rect(x, 192, 888, 80, tint, radius=22, name=prefix + "-标题底色")
    s.rect(x, 236, 888, 36, tint, name=prefix + "-标题底色平底")
    s.ellipse(x + 43, 232, 23, 23, accent, name=prefix + "-编号圆")
    s.text(x + 43, 240, number, 21, C["white"], True, "middle",
           name=prefix + "-编号文字")
    s.text(x + 82, 245, heading, 34, accent, True, width=770,
           name=prefix + "-问题标题")
    s.text(x + 30, 314, question, 28, C["ink"], True, width=826,
           name=prefix + "-核心问句")


def draw_bulk(s):
    # Liquid reference: uninterrupted schematic liquid channels, not a measured
    # tortuosity or a claim that ions in real liquids follow straight tracks.
    s.rect(78, 352, 188, 365, "#F0F9FD", "#AACFE4", 1.5, 12,
           name="问题一-液态参照外框")
    s.rect(78, 352, 188, 53, "#E2F1FA", radius=12, name="问题一-液态标题底色")
    s.rect(78, 379, 188, 26, "#E2F1FA", name="问题一-液态标题底色平底")
    s.text(172, 389, "液态电解液", 25, C["blue_dark"], True, "middle",
           name="问题一-液态标题")
    for j, x in enumerate((129, 211), 1):
        s.rect(x - 13, 420, 26, 242, "#D9EFF9", radius=13,
               name=f"问题一-液态-连续液相通道{j}")
        s.arrow(f"M {x} 650 L {x} 432", x, 432, -90, C["blue"], 4.3,
                "7 6", name=f"问题一-液态-迁移箭头{j}")
    for j, (x, y) in enumerate(((129, 541), (211, 590)), 1):
        s.ellipse(x, y, 29, 29, "#F6FCFF", "#76BAD9", 1.2, "4 4",
                  name=f"问题一-液态-离子环境{j}")
        for k, (dx, dy) in enumerate(((-20, -18), (21, -15), (1, 23)), 1):
            s.ellipse(x + dx, y + dy, 6, 6, C["solvent"],
                      name=f"问题一-液态-溶剂{j}-{k}")
        s.li(x, y, 17, f"问题一-液态-锂离子{j}")
    s.line(91, 674, 253, 674, "#B9D7E8", 1.2, name="问题一-液态-图注分隔线")
    s.text(172, 701, "连续液相通道", 21, C["blue_dark"], True, "middle",
           name="问题一-液态-图注")

    # Gel: the drawing illustrates possible pathway / coordination limitations;
    # it deliberately does not assert that every network causes these effects.
    s.rect(292, 352, 614, 365, "#FBFCFE", C["edge"], 1.5, 12,
           name="问题一-凝胶示意外框")
    s.rect(292, 352, 614, 53, "#EEF2F7", radius=12,
           name="问题一-凝胶标题底色")
    s.rect(292, 379, 614, 26, "#EEF2F7", name="问题一-凝胶标题底色平底")
    s.text(599, 389, "传统凝胶：聚合物骨架 + 溶剂", 25, C["ink"], True, "middle",
           width=574, name="问题一-凝胶标题")
    s.ellipse(546, 508, 54, 54, "#FEF0EE", "#DDB2A9", 1.3, "5 5",
              name="问题一-凝胶-局部配位滞留区域")
    paths = (
        "M 315 443 C 371 426 413 457 475 441 C 534 425 582 452 637 438 C 708 424 783 464 886 438",
        "M 312 554 C 374 531 415 570 485 547 C 547 526 599 564 666 546 C 741 527 800 562 884 543",
        "M 312 642 C 378 624 423 651 488 634 C 554 617 609 648 675 632 C 741 616 808 650 884 633",
        "M 382 420 C 364 461 399 487 380 530 C 359 579 394 611 382 658",
        "M 594 419 C 567 469 608 489 589 535 C 569 583 604 611 591 659",
        "M 791 419 C 764 461 805 492 785 536 C 764 579 799 612 784 659",
    )
    for j, d in enumerate(paths, 1):
        s.path(d, C["mesh"], 6.2, name=f"问题一-凝胶-聚合物链{j}")
    for j, (x, y) in enumerate(((382, 441), (590, 440), (787, 442),
                                (380, 547), (588, 548), (786, 548),
                                (384, 634), (592, 634), (784, 634)), 1):
        node(s, x, y, f"问题一-凝胶-节点{j}")

    for j, (x, y) in enumerate(((331, 479), (431, 605), (644, 602),
                                (839, 590), (687, 489), (870, 486)), 1):
        solvent(s, x, y, f"问题一-凝胶-自由溶剂{j}")

    s.ellipse(772, 523, 75, 45, "#EEF1F5", "#9AA9B8", 1.7, "5 5",
              name="问题一-凝胶-局部团聚区域")
    s.text(772, 520, "局部团聚", 22, C["muted"], True, "middle",
           name="问题一-凝胶-团聚文字")
    s.text(772, 546, "通道不连续", 18, C["muted"], False, "middle",
           name="问题一-凝胶-通道文字")

    s.arrow("M 331 651 C 355 626 367 591 349 565 C 328 532 363 509 405 503 "
            "C 454 496 447 425 489 418 C 538 410 553 430 608 442 "
            "C 668 456 647 516 675 544 C 697 570 725 590 752 577 "
            "C 807 553 844 487 873 426",
            873, 426, -63, C["gold"], 4.5, "7 6", name="问题一-凝胶-曲折迁移路径")
    s.text(734, 664, "更曲折的路径", 22, C["gold"], True, "middle",
           name="问题一-凝胶-曲折路径标注")

    # Generic polar site, without implying that a specific functional group is
    # common to every conventional polymer chemistry.
    s.line(501, 445, 515, 466, C["mesh"], 3, name="问题一-凝胶-极性位点连接")
    s.ellipse(520, 473, 9, 9, C["coral"], name="问题一-凝胶-极性位点")
    s.line(525, 482, 538, 496, C["coral"], 2.2, "4 4",
           name="问题一-凝胶-位点配位虚线")
    for j, (x, y) in enumerate(((579, 482), (581, 533), (511, 535)), 1):
        solvent(s, x, y, f"问题一-凝胶-局部溶剂{j}", 13)
    s.li(546, 508, 21, "问题一-凝胶-滞留锂离子")
    s.li(393, 509, 16, "问题一-凝胶-迁移锂离子1")
    s.li(837, 493, 16, "问题一-凝胶-迁移锂离子2")
    s.pill(551, 595, 193, 37, "局部配位滞留", "#FFF6F3", "#DDAA9E",
           C["coral"], 22, name="问题一-凝胶-配位滞留标注")
    s.line(550, 576, 548, 552, "#C47D6B", 1.7, "4 4",
           name="问题一-凝胶-配位标注引线")
    s.line(307, 674, 891, 674, C["edge"], 1.2, name="问题一-凝胶-图注分隔线")
    s.text(599, 701, "路径曲折 + 局部滞留", 24, C["blue_dark"], True, "middle",
           name="问题一-凝胶-图注")


def draw_interface(s):
    s.rect(1014, 352, 828, 365, "#FBFCFE", C["edge"], 1.5, 12,
           name="问题二-界面示意外框")
    s.text(1150, 392, "CEI 不均", 22, C["blue_dark"], True, "middle",
           name="问题二-正极-膜标签")
    s.text(1428, 392, "传统凝胶", 25, C["ink"], True, "middle",
           name="问题二-凝胶标签")
    s.text(1705, 392, "SEI 不均", 22, C["coral"], True, "middle",
           name="问题二-负极-膜标签")

    for x, label, color, prefix in ((1058, "正极", "#73839A", "问题二-正极"),
                                     (1770, "负极", "#B7A18C", "问题二-负极")):
        s.rect(x, 421, 34, 246, color, radius=4, name=prefix + "-电极主体")
        s.rect(x + 3, 424, 5, 240, "#9CABB9" if label == "正极" else "#D4C1AD",
               radius=2, name=prefix + "-电极高光")
        s.text(x + 17, 701, label, 22, C["muted"], True, "middle",
               name=prefix + "-电极文字")

    # Irregular film blocks are qualitative cartoons; no chemistry or fraction
    # is assigned to the colors, and neither electrode is specified as metal.
    for j, (y, h, width) in enumerate(((423, 34, 30), (472, 22, 20),
                                      (507, 35, 34), (557, 20, 23),
                                      (624, 40, 29)), 1):
        s.rect(1097, y, width, h, "#A9C7D9" if j % 2 else "#CBDCE8", radius=6,
               name=f"问题二-正极-不均匀界面膜{j}")
    for j, (y, h, width) in enumerate(((423, 35, 24), (472, 25, 33),
                                      (512, 38, 26), (565, 30, 30),
                                      (629, 36, 23)), 1):
        s.rect(1766 - width, y, width, h, "#DFC2B4" if j % 2 else "#EBD7CC", radius=6,
               name=f"问题二-负极-不均匀界面膜{j}")

    s.ellipse(1117, 601, 24, 22, C["white"], C["coral"], 2, "5 4",
              name="问题二-正极-局部接触空隙")
    s.path("M 1742 571 L 1754 578 L 1745 586 L 1760 593", C["coral"], 2.7,
           name="问题二-负极-膜层裂纹")

    paths = (
        "M 1231 447 C 1314 428 1383 463 1452 446 C 1515 433 1564 453 1628 445",
        "M 1227 547 C 1305 527 1377 562 1453 544 C 1515 527 1578 555 1631 542",
        "M 1228 632 C 1313 613 1386 646 1458 626 C 1521 608 1578 638 1633 627",
        "M 1281 422 C 1261 464 1297 491 1279 533 C 1260 580 1293 614 1280 655",
        "M 1442 421 C 1422 465 1458 492 1440 535 C 1421 579 1454 614 1442 655",
        "M 1597 422 C 1577 465 1613 492 1595 533 C 1576 579 1609 613 1597 653",
    )
    for j, d in enumerate(paths, 1):
        s.path(d, C["mesh"], 5.6, name=f"问题二-凝胶-聚合物链{j}")
    for j, (x, y) in enumerate(((1280, 443), (1440, 447), (1597, 447),
                                (1278, 544), (1440, 545), (1595, 546)), 1):
        node(s, x, y, f"问题二-凝胶-节点{j}", 6)
    for j, (x, y) in enumerate(((1343, 484), (1487, 489), (1359, 591), (1521, 602)), 1):
        solvent(s, x, y, f"问题二-凝胶-溶剂{j}", 13)
    for j, (x, y, r) in enumerate(((1409, 511, 21), (1251, 584, 16),
                                   (1518, 462, 16)), 1):
        s.li(x, y, r, f"问题二-凝胶-锂离子{j}")

    # Different arrow weights/spacings illustrate a possible uneven local flux,
    # not a measured energy barrier or deposition profile.
    s.arrow("M 1450 511 C 1533 486 1599 492 1709 510", 1709, 510, 8,
            C["blue"], 5, "7 5", name="问题二-负极-较高局部离子通量")
    s.arrow("M 1471 613 C 1565 630 1623 608 1711 618", 1711, 618, 5,
            "#85B6D4", 2.4, "5 6", name="问题二-负极-较低局部离子通量")
    s.arrow("M 1242 496 C 1208 480 1175 484 1149 502", 1149, 502, 154,
            C["blue"], 3.5, "6 5", name="问题二-正极-界面迁移箭头")
    s.pill(1282, 686, 216, 36, "局部接触空隙", "#FFF6F3", "#DDAA9E",
           C["coral"], 22, name="问题二-正极-空隙标注")
    s.path("M 1180 677 L 1143 637 L 1133 614", C["coral"], 1.7, dash="4 4",
           name="问题二-正极-空隙标注引线")
    s.pill(1568, 576, 216, 36, "反应膜破裂 / 重构", "#FFF6F3", "#DDAA9E",
           C["coral"], 21, name="问题二-负极-成膜失稳标注")
    s.path("M 1676 576 L 1702 573 L 1733 582", C["coral"], 1.7, dash="4 4",
           name="问题二-负极-成膜标注引线")
    s.text(1631, 685, "局部通量不均", 22, C["blue_dark"], True, "middle",
           name="问题二-负极-通量标注")


def bullets(s, x, items, color, prefix):
    for j, (title, text) in enumerate(items, 1):
        y = 799 + (j - 1) * 43
        s.ellipse(x + 6, y - 9, 4.5, 4.5, color, name=f"{prefix}-要点{j}-圆点")
        s.text(x + 25, y, title + text, 24, C["ink"], False, width=792,
               name=f"{prefix}-要点{j}-文字")


def build(font_dir):
    s = Scene(font_dir=font_dir)
    s.rect(0, 0, 1920, 1080, C["bg"], name="页面-背景")
    s.rect(48, 41, 7, 98, C["blue"], radius=3, name="页面-标题强调线")
    s.text(76, 61, "SCIENTIFIC QUESTIONS", 20, C["blue"], True, name="页面-英文页眉")
    s.text(76, 119, TITLE, 45, C["ink"], True, width=1795, name="页面-主标题")
    s.text(76, 162, "仅讨论传统凝胶的常见瓶颈：体相传输与电极界面", 26, C["muted"],
           name="页面-副标题")

    panel(s, 48, "01", "体相：离子传输受限",
          "为什么引入聚合物网络后，电导率可能下降？", C["blue"], C["blue_light"], "问题一")
    panel(s, 984, "02", "界面：接触与成膜失稳",
          "为什么接触与成膜不均会导致界面阻抗增长？", C["coral"], C["coral_light"], "问题二")
    draw_bulk(s)
    draw_interface(s)

    s.text(492, 748, "网络限域 → 有效离子迁移可能受限", 26, C["blue_dark"], True, "middle",
           width=828, name="问题一-宏观总结")
    s.text(1428, 748, "接触不均 + 界面膜失稳 → 阻抗增长", 26, C["coral"], True, "middle",
           width=828, name="问题二-宏观总结")
    s.line(77, 769, 907, 769, C["edge"], 1.2, name="问题一-要点分隔线")
    s.line(1013, 769, 1843, 769, C["edge"], 1.2, name="问题二-要点分隔线")
    bullets(s, 78, (
        ("通道曲折：", "网络分割连续液相，离子需要绕行"),
        ("配位 / 链段约束：", "局部结合与松弛可能限制迁移"),
        ("结构不均：", "团聚或相分离使有效通道不连续"),
    ), C["blue"], "问题一")
    bullets(s, 1014, (
        ("接触不均：", "局部空隙减少有效电极接触面积"),
        ("界面副反应：", "SEI / CEI 不均、破裂与重构"),
        ("通量失衡：", "局部极化增大，长期阻抗持续增长"),
    ), C["coral"], "问题二")

    s.rect(48, 946, 1824, 77, "#172F4D", radius=16, name="过渡-底部问题栏")
    s.pill(180, 985, 216, 40, "引出体系设计", "#2D698A", "#2D698A", C["white"], 23,
           name="过渡-提示标签")
    s.text(318, 995, "如何兼顾稳定的聚合物骨架、高效离子传输与长期稳定的电极界面？",
           29, C["white"], True, width=1515, name="过渡-核心设计目标问句")
    s.text(48, 1055, "注：图示为常见挑战的定性示意，并非所有凝胶必然失效；迁移数不必然下降，影响程度依网络、溶剂与电极体系而变。",
           19, C["muted"], width=1818, name="页面-科学边界说明")

    s.validate()
    names = [e.name for e in s.elements]
    assert len(names) == len(set(names)), "Every editable object needs a unique name."
    svg = s.svg(title=TITLE, description=DESCRIPTION)
    ET.fromstring(svg)
    for forbidden in ("MBA", "DMTFA", "TTE", "LiDFOB"):
        assert forbidden not in svg
    return s


NOTES = """【单页：传统凝胶电解质的两大共性挑战】
本页只介绍传统体系的常见宏观问题，不展示本工作配方或具体设计方案。

【约一分钟汇报讲稿】
这一页先总结凝胶电解质可能面临的两类挑战。
第一是体相离子传输。液态电解液具有连续液相；引入聚合物网络后，网络限域、局部配位和链段松弛可能改变离子的有效迁移路径。若骨架团聚或网络与溶剂相容不足，还可能出现局部通道不连续。因此需要关注保液、网络稳定性与高效传输之间的平衡，而不是只追求更高的交联度。
第二是电极界面。局部接触不充分、界面膜不均以及反复破裂和重构，可能引起局部离子通量不均、持续副反应、极化及阻抗增长。
由此提出下一页的设计目标：怎样在形成稳定骨架的同时，维持高效离子传输和长期稳定的电极界面？

【科学表述边界】
1. 共性挑战是常见问题类别，不表示每一种传统凝胶都存在图中所有机制。离子电导率不必然下降；锂离子迁移数也不必然下降，两者不能混为同一结论。
2. 网络图、液相路径和溶剂符号均为定性示意，不代表实际孔径、曲折度、配位数或溶剂化结构。
3. 极性位点表示候选局部相互作用，不指定所有聚合物都含同一官能团，也不意味着一切聚合物配位都阻碍迁移。
4. 界面图使用通用正极/负极；没有指定锂金属或某种正极材料，因此不把枝晶当作所有电极必然发生的现象。
5. 凝胶含有液相，界面不能简单一概称为固–固接触。原位成胶或合适的配方也可能改善接触与润湿。
6. 界面膜色块没有指定具体有机/无机组分、组成比例、膜厚或能垒，不是显微照片、计算结果或化学组分证据。

【逐个编辑方法】
该文件仅一页16:9；标题、文字、骨架、节点、溶剂符号、离子、虚线、箭头、电极和膜块都是PowerPoint原生对象，没有嵌入图片。
在“开始 → 选择 → 选择窗格”中按“问题一-”“问题二-”“页面-”“过渡-”等唯一对象名定位。文字可双击编辑；骨架等自由曲线可右键“编辑顶点”；颜色、线宽和位置均可修改。
PNG仅作高清预览；SVG是额外的矢量源，并未作为图片嵌入PPT。
"""


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--font-dir", type=Path, default=Path("/tmp/fonts"))
    args = parser.parse_args()
    scene = build(args.font_dir)
    svg = scene.svg(title=TITLE, description=DESCRIPTION)
    BASE.with_suffix(".svg").write_text(svg, encoding="utf-8")
    BASE.with_suffix(".png").write_bytes(resvg_py.svg_to_bytes(
        svg_string=svg, width=3840, height=2160, font_dirs=[str(args.font_dir)]
    ))
    prs = scene.powerpoint(
        BASE.with_suffix(".pptx"), NOTES, title=TITLE,
        subject="Conventional gel polymer electrolyte: bulk transport and interphase challenges",
        keywords="GPE, conventional gel, scientific questions, fully editable native shapes",
    )
    assert len(prs.slides) == 1
    assert all(shape.shape_type != MSO_SHAPE_TYPE.PICTURE for shape in prs.slides[0].shapes)
    bundle = HERE / (BASE.name + "_Files.zip")
    with ZipFile(bundle, "w", ZIP_DEFLATED, compresslevel=9) as archive:
        for suffix in (".pptx", ".svg", ".png"):
            path = BASE.with_suffix(suffix)
            archive.write(path, path.name)
        guide = HERE / (BASE.name + "_Notes.md")
        archive.write(guide, "README_EDITING.md")
    print(f"Created single-slide editable scientific question page: {len(scene.elements)} native objects.")
    print(BASE.with_suffix(".pptx"))


if __name__ == "__main__":
    main()
