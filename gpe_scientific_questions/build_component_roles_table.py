#!/usr/bin/env python3
"""Create a concise, natively editable PPT table of GPE component roles.

The PowerPoint uses a real 7×3 table: every cell is editable. The preview SVG
is independently drawn from the same content, never embedded as a PPT image.
No composition or molecular weight is invented for the PEGDA control.
"""
from __future__ import annotations

import argparse
import csv
from pathlib import Path

import resvg_py
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

from build_design_concept import Scene, font_runs

HERE = Path(__file__).resolve().parent
BASE = HERE / "GPE_Component_Roles_Table"
TITLE = "体系各组分的作用与选择依据"
FORMULA = "MBA 交联网络  |  1.0 M LiPF₆ + 0.2 M LiDFOB  |  DMTFA : TTE = 1 : 1"
CONTROL = "PEGDA 为对比样的网络基体，不是本体系同时加入的组分"
HEADERS = ("组分 / 定位", "主要作用", "选择理由")
# label, subtitle, role, rationale
ROWS = (
    ("MBA", "交联单体",
     "双键构网，固定电解液；酰胺位点参与分子间作用。",
     "研究酰胺网络对保液、相容性与离子传输的调控。"),
    ("LiPF₆（1.0 M）", "主体锂盐",
     "提供 Li⁺，构成电解液的基础导电组分。",
     "成熟常用锂盐，作为可比较的基础电解液。"),
    ("LiDFOB（0.2 M）", "成膜调控共盐",
     "提供 DFOB⁻；其分解产物可参与电极界面膜形成。",
     "含 F/B，可探索富 LiF / B–O 的 SEI / CEI 成膜。"),
    ("DMTFA", "主溶剂",
     "作为主要溶盐 / 配位组分；羰基可与 Li⁺ 相互作用。",
     "与 MBA 同含酰胺结构，可探索网络相容与配位调控。"),
    ("TTE", "弱配位氟醚稀释组分",
     "调节混合液相环境，不作为主要溶盐组分。",
     "与 DMTFA 搭配，调节配位环境和液相性质。"),
    ("PEGDA（对比样）", "聚醚型交联基体",
     "两端双键构网；柔性聚醚链可与 Li⁺ 相互作用。",
     "常见凝胶基体、同为双官能度；对比酰胺网络与聚醚网络。"),
)
FOOTERS = (
    "公平对比：建议液相配方与测试条件一致；注明 PEGDA 分子量、聚合物含量，并评估实际成胶 / 交联程度。",
    "表中为功能定位与设计理由，具体相容、配位及成膜效果需验证；仅凭 1:1 配比不能认定 LHCE。",
)
X, Y = 48, 246
WIDTHS = (350, 728, 746)
HEIGHTS = (74,) + (100,) * len(ROWS)
INK, BLUE, LINE = "#172B46", "#285E8B", "#DCE4ED"


def wrap(scene, text, size, width):
    lines, line = [], ""
    for char in text:
        if line and scene.measure(line + char, size) > width:
            lines.append(line.rstrip())
            line = char
        else:
            line += char
    if line:
        lines.append(line.rstrip())
    return lines


def decorations(font_dir):
    scene = Scene(font_dir=font_dir)
    scene.rect(0, 0, 1920, 1080, "#F5F7FB", name="表格-页面背景")
    scene.rect(48, 40, 7, 99, "#285E8B", radius=3, name="表格-标题强调线")
    scene.text(76, 62, "COMPONENT ROLES & CONTROL", 20, BLUE, True, name="表格-英文页眉")
    scene.text(76, 122, TITLE, 47, INK, True, name="表格-主标题")
    scene.text(76, 168, FORMULA, 27, BLUE, True, width=1780, name="表格-体系配方")
    scene.text(76, 211, CONTROL, 24, "#63788F", name="表格-对比样说明")
    for j, text in enumerate(FOOTERS):
        scene.text(48, 976 + j * 43, text, 21, "#63788F", width=1824,
                   name=f"表格-说明{j+1}")
    scene.validate()
    return scene


def preview(scene):
    positions = (X, X + WIDTHS[0], X + WIDTHS[0] + WIDTHS[1])
    for col, (x, w, title) in enumerate(zip(positions, WIDTHS, HEADERS)):
        scene.rect(x, Y, w, HEIGHTS[0], "#213D5A", "#213D5A", 1,
                   name=f"预览-表头底色{col}")
        scene.text(x + 24, Y + 48, title, 28, "#FFFFFF", True,
                   name=f"预览-表头文字{col}")
    for row, (label, subtitle, role, rationale) in enumerate(ROWS):
        y = Y + HEIGHTS[0] + row * HEIGHTS[row+1]
        fill = "#EEF5FB" if row == 5 else "#FFFFFF" if row % 2 == 0 else "#F6F9FC"
        for col, (x, w) in enumerate(zip(positions, WIDTHS)):
            scene.rect(x, y, w, HEIGHTS[row+1], fill, LINE, 1.1,
                       name=f"预览-单元格{row}-{col}")
        scene.text(X + 24, y + 43, label, 29, BLUE, True,
                   width=WIDTHS[0]-48, name=f"预览-组分名称{row}")
        scene.text(X + 24, y + 77, subtitle, 22, "#63788F",
                   width=WIDTHS[0]-48, name=f"预览-组分定位{row}")
        for col, text in ((1, role), (2, rationale)):
            lines = wrap(scene, text, 26, WIDTHS[col] - 48)
            assert len(lines) <= 2, (label, lines)
            baseline = y + 50 - (len(lines)-1)*16
            for j, line in enumerate(lines):
                scene.text(positions[col] + 24, baseline + j * 34, line, 26, INK,
                           name=f"预览-正文{row}-{col}-{j}")
    scene.validate()
    return scene


def set_text(cell, lines, k, sizes=None, bold=False, color=INK):
    cell.margin_left = cell.margin_right = int(round(24*k))
    cell.margin_top = cell.margin_bottom = int(round(7*k))
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame
    tf.word_wrap = False  # preview/PPT have the same explicit line breaks
    tf.clear()
    for j, text in enumerate(lines):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph()
        p.space_before = p.space_after = Pt(0)
        p.line_spacing = 1.12
        size = sizes[j] if sizes else 26
        for chinese, part in font_runs(text):
            run = p.add_run()
            run.text = part
            run.font.name = "Microsoft YaHei" if chinese else "Arial"
            run.font.size = Pt(size/2)
            run.font.bold = bold if sizes is None or j == 0 else False
            run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
            if chinese:
                ea = OxmlElement("a:ea")
                ea.set("typeface", "Microsoft YaHei")
                run._r.get_or_add_rPr().append(ea)


def cell_fill_and_border(cell, fill, border=LINE):
    cell.fill.solid()
    cell.fill.fore_color.rgb = RGBColor.from_string(fill.lstrip("#"))
    tcpr = cell._tc.get_or_add_tcPr()
    for index, side in enumerate(("L", "R", "T", "B")):
        ln = OxmlElement("a:ln" + side)
        ln.set("w", "6350")
        solid = OxmlElement("a:solidFill")
        col = OxmlElement("a:srgbClr")
        col.set("val", border.lstrip("#"))
        solid.append(col)
        ln.append(solid)
        tcpr.insert(index, ln)


NOTES = """【各组分作用与PEGDA对比理由】
配方沿用已给出的MBA、1.0 M LiPF6 + 0.2 M LiDFOB、DMTFA:TTE=1:1。
不自行补充MBA含量、成胶条件、引发剂、溶剂比基准或PEGDA对比样配方。
PEGDA是聚乙二醇二丙烯酸酯，是对比样的交联基体，不是本体系额外加入的组分。

【PEGDA为何适合作为对比】
已有锂电池凝胶电解质研究采用PEGDA作为交联基体；其两个可聚合端基与MBA的双官能度构网在设计层面具有可比性。
PEGDA提供柔性聚醚链，而MBA形成含酰胺的网络，适合比较网络化学、链段特性及其与液态电解液相互作用的整体差异。
这不意味着PEGDA性能较差，也不能将对比结果自动归因于单一氢键或官能团作用。
建议保持液态电解液与测试条件一致；注明PEGDA分子量、聚合物含量、转化/凝胶分数和保液程度。双官能度相同、等质量加入均不等于实际交联密度相同。

【科学边界】
酰胺结构相近不证明一定相容。候选氢键供体是MBA的N–H，DMTFA羰基是候选受体；DMTFA本身无N–H。
DFOB阴离子的含F/B结构提供成膜调控动机，但当前正负极真实膜组分、优先反应次序和性能改善仍需证据。
TTE作为弱配位氟醚稀释组分的设计定位，不代表本体系已经形成LHCE、阴离子必然进入第一配位鞘或电导率必然提升。

【PEGDA背景文献；不作为本配方的性能证据】
[2](https://doi.org/10.3390/gels9120975) Gel Polymer Electrolytes for Lithium-Ion Batteries Enabled by Photo Crosslinked Polymer Network, Gels, 2023. 文献使用PEGDA/DPHA，不直接等同本PEGDA对照。
[3](https://www.sciencedirect.com/science/article/abs/pii/S0013468613025176) Enhanced separator properties by thermal curing of poly(ethylene glycol)diacrylate-based gel polymer electrolytes for lithium-ion batteries. 文献采用碳酸酯电解液，不直接等同本DMTFA/TTE体系。

【编辑说明】
PPT含一张原生7×3表格；可直接编辑所有单元格、调整行高列宽和单元格颜色。预览PNG没有嵌入PPT。
"""


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--font-dir", type=Path, default=Path("/tmp/fonts"))
    args = parser.parse_args()
    scene = decorations(args.font_dir)
    prs = scene.powerpoint(BASE.with_suffix(".pptx"), NOTES,
        title=TITLE, subject="GPE component roles and PEGDA control rationale",
        keywords="MBA, LiPF6, LiDFOB, DMTFA, TTE, PEGDA, native editable table")
    k = prs.slide_width / scene.width
    emu = lambda v: int(round(v*k))
    frame = prs.slides[0].shapes.add_table(7, 3, emu(X), emu(Y), emu(sum(WIDTHS)), emu(sum(HEIGHTS)))
    frame.name = "组分作用-原生可编辑表格"
    table = frame.table
    table.first_row = table.last_row = table.first_col = table.last_col = False
    table.horz_banding = table.vert_banding = False
    for col, w in enumerate(WIDTHS):
        table.columns[col].width = emu(w)
    for row, h in enumerate(HEIGHTS):
        table.rows[row].height = emu(h)
    for col, text in enumerate(HEADERS):
        cell = table.cell(0, col)
        cell_fill_and_border(cell, "#213D5A", "#213D5A")
        set_text(cell, [text], k, [28], True, "#FFFFFF")
    for row, (label, subtitle, role, rationale) in enumerate(ROWS, 1):
        fill = "#EEF5FB" if row == 6 else "#FFFFFF" if row % 2 else "#F6F9FC"
        for col in range(3):
            cell_fill_and_border(table.cell(row, col), fill)
        set_text(table.cell(row, 0), [label, subtitle], k, [29, 22], True, BLUE)
        for col, text in ((1, role), (2, rationale)):
            set_text(table.cell(row, col), wrap(scene, text, 26, WIDTHS[col]-48), k)
    prs.save(BASE.with_suffix(".pptx"))

    full = preview(scene)
    svg = full.svg(title=TITLE, description="Component roles and PEGDA comparator rationale; mechanistic benefits require validation.")
    BASE.with_suffix(".svg").write_text(svg, encoding="utf-8")
    BASE.with_suffix(".png").write_bytes(resvg_py.svg_to_bytes(
        svg_string=svg, width=3840, height=2160, font_dirs=[str(args.font_dir)]))
    with BASE.with_suffix(".csv").open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.writer(f, lineterminator="\n")
        writer.writerow(HEADERS)
        writer.writerows((label+"（"+subtitle+"）", role, why) for label, subtitle, role, why in ROWS)
    md = ["# " + TITLE, "", FORMULA, "", CONTROL, "",
          "| " + " | ".join(HEADERS) + " |", "| --- | --- | --- |"]
    md += ["| " + " | ".join((label+" · "+subtitle, role, why)) + " |" for label, subtitle, role, why in ROWS]
    md += ["", "## 对比与表述说明", ""] + list(FOOTERS)
    md += ["", "PEGDA：聚乙二醇二丙烯酸酯；双官能度相同不等于交联密度相同。", "",
           "PEGDA应用背景：[2](https://doi.org/10.3390/gels9120975)；[3](https://www.sciencedirect.com/science/article/abs/pii/S0013468613025176)。文献配方不同，不作为当前体系的性能证据。",
           "", "PPT为7×3原生表格，可直接编辑单元格；CSV可在Excel中打开。PNG仅用于预览。"]
    BASE.with_suffix(".md").write_text("\n".join(md)+"\n", encoding="utf-8")
    print("Created editable native table: 7 rows × 3 columns, 1 widescreen slide.")
    print(BASE.with_suffix(".pptx"))


if __name__ == "__main__":
    main()
