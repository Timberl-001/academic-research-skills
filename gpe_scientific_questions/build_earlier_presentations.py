#!/usr/bin/env python3
"""Package the earlier original GPE figures and existing PowerPoint decks.

Unlike the latest design-concept slide, these historical figures are embedded
PNGs, not natively editable molecular drawings. Titles and cautions added by
this builder are native text boxes. Editable SVG sources accompany the bundle.

Run from any directory:
    python3 gpe_scientific_questions/build_earlier_presentations.py
"""
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
DECK = HERE / "MBA_GPE_Earlier_Schematics_Collection.pptx"
BUNDLE = HERE / "GPE_PPT_Collection.zip"
FIGURES = (
    ("早期方案一：三项科学问题对比矩阵", "MBA_GPE_Scientific_Questions_Comparison"),
    ("早期方案二：全电池与分子机制示意", "MBA_GPE_FullCell_Molecular_Mechanism"),
    ("早期方案三（左图）：凝胶化与离子传输", "Fig1_Why_Conductivity_Drops_After_Gelation"),
    ("早期方案三（右图）：脱溶剂化与电极界面", "Fig2_Interface_Desolvation_and_SEI_Issue"),
    ("早期方案四：双图布局（包含本体系）", "MBA_GPE_Two_Pictures_PPT_Slide"),
)
CAUTION = "早期概念草稿，仅供回看；配位、动力学与成膜机理均需验证，请以最新设计页的“设计预期”表述为准。"
NOTES = """【历史方案归档，不作为已验证的体系机理】
图件是此前按讨论过程绘制的概念版本，未复制任何参考文献图像。
旧图部分措辞较强，保存仅便于回看布局与设计演变，不构成实验结果。
仅凭1.0 M LiPF6 + 0.2 M LiDFOB和DMTFA:TTE=1:1不能认定LHCE、CIP/AGG占优，
不能证明阴离子进入第一溶剂化鞘、降低脱溶剂化能垒、锚定阴离子、提高迁移数或清除HF。
右侧配位、传输与LiF/B–O成膜均应理解为待验证的设计假设，不能作为已测得的结论汇报。
正式体系设计汇报优先使用MBA_GPE_Design_Concept_Slide.pptx，并参照其中的备注。
这套早期合集的图为嵌入PNG，不能在PPT内逐个编辑图中元素；原生标题/提示可编辑。
压缩包的SVG_sources目录附原始矢量源，可在支持SVG编辑的软件中修改旧图。
"""
SVG_STEMS = [stem for _, stem in FIGURES] + [
    "Common_Q1_Bulk_Conductivity_Drop",
    "Common_Q2_Interface_Degradation",
    "Traditional_GPE_Two_Common_Problems_Slide",
    "MBA_GPE_Design_Concept_Slide",
    "MBA_GPE_Design_Comparison_Figure",
]


def textbox(slide, text, x, y, w, h, size, color, bold=False, name=""):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name or text[:40]
    frame = box.text_frame
    frame.margin_left = frame.margin_right = 0
    frame.margin_top = frame.margin_bottom = 0
    frame.word_wrap = True
    run = frame.paragraphs[0].add_run()
    run.text = text
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    return box


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.3333333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "MBA凝胶课题：早期示意图方案归档"
    prs.core_properties.subject = "Historical conceptual drafts, not experimentally verified mechanisms"
    for i, (title, stem) in enumerate(FIGURES, 1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F6F8FB")
        header = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.62)
        )
        header.name = "可编辑页眉底色"
        header.fill.solid()
        header.fill.fore_color.rgb = RGBColor.from_string("172B46")
        header.line.fill.background()
        textbox(slide, f"{i:02d}  {title}", 0.28, 0.11, 12.7, 0.39,
                20, "FFFFFF", True, "可编辑标题")
        with Image.open(HERE / f"{stem}.png") as image:
            width, height = image.size
        scale = min(12.78 / width, 6.02 / height)
        w, h = width * scale, height * scale
        picture = slide.shapes.add_picture(
            str(HERE / f"{stem}.png"),
            Inches((13.3333333333 - w) / 2), Inches(0.79 + (6.02 - h) / 2),
            width=Inches(w), height=Inches(h),
        )
        picture.name = f"历史图件PNG（矢量源见SVG_sources/{stem}.svg）"
        textbox(slide, CAUTION, 0.3, 7.01, 12.7, 0.32, 10.5,
                "9D5A48", name="可编辑历史草稿说明")
        slide.notes_slide.notes_text_frame.text = f"{title}\n\n{NOTES}"
    prs.save(DECK)


def build_bundle():
    files = [
        (HERE / "Traditional_GPE_Scientific_Questions_Slide.pptx",
         "01_Traditional_GPE_Scientific_Questions_Slide.pptx"),
        (DECK, "02_MBA_GPE_Earlier_Schematics_Collection.pptx"),
        (HERE / "MBA_GPE_Design_Concept_Slide.pptx",
         "03_MBA_GPE_Design_Concept_Slide.pptx"),
        (HERE / "PPT_File_Index.md", "README_PPT_File_Index.md"),
    ]
    files += [(HERE / f"{stem}.svg", f"SVG_sources/{stem}.svg") for stem in SVG_STEMS]
    for path, _ in files:
        if not path.exists():
            raise FileNotFoundError(path)
    with ZipFile(BUNDLE, "w", ZIP_DEFLATED, compresslevel=9) as archive:
        for path, name in files:
            archive.write(path, name)
    with ZipFile(BUNDLE) as archive:
        assert archive.testzip() is None
    print(f"Created {DECK.name}: {len(FIGURES)} historical draft slides")
    print(f"Created {BUNDLE.name}: 3 distinct PPT decks + 10 editable SVG sources + index")


if __name__ == "__main__":
    build_deck()
    build_bundle()
