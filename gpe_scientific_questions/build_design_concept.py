#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Original, editable MBA GPE design-concept slide.

The figure uses the user's reference only for its side-by-side narrative, not
for its image pixels, molecular systems, electrode names, or empirical claims.
Every visible element is drawn here.  SVG and PPTX share the same scene graph.

Usage:
    python3 gpe_scientific_questions/build_design_concept.py --font-dir /tmp/fonts
Dependencies: pillow, python-pptx, resvg-py
"""
from __future__ import annotations

import argparse
import math
import re
import xml.etree.ElementTree as ET
from dataclasses import dataclass
from html import escape
from pathlib import Path
from typing import Any

import resvg_py
from PIL import ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

OUT = Path(__file__).resolve().parent
W, H = 1920, 1080
P = {
    'bg': '#F7F9FC', 'white': '#FFFFFF', 'ink': '#172B46',
    'muted': '#68798D', 'line': '#DFE7EF', 'blue': '#497FB5',
    'blue_light': '#EFF5FB', 'blue_mid': '#D7E7F7',
    'teal': '#148F84', 'teal_dark': '#087469', 'mint': '#7DCABB',
    'teal_light': '#E7F5F1', 'coral': '#D67661',
    'coral_light': '#FCF0EB', 'gold': '#E9B943', 'li': '#EBAE3E',
    'li_dark': '#C28723', 'solvent': '#F6DCD0', 'solvent_line': '#D69985',
    'tte': '#DCEFEF', 'tte_line': '#8EBBBE', 'pf': '#A69DCB',
    'dfob': '#62AD98', 'organic': '#B7CDDF', 'organic2': '#D3DDF0',
    'poly_old': '#8EA5B9', 'electrode_c': '#6E7D91',
    'electrode_a': '#C6AD94',
}


def rgb(h: str):
    return RGBColor.from_string(h.lstrip('#').upper())


def han(c: str):
    n = ord(c)
    return (0x2E80 <= n <= 0x9FFF) or (0xF900 <= n <= 0xFAFF) or (0xFF00 <= n <= 0xFFEF)


def font_runs(s: str):
    groups = []
    for c in s:
        chinese = han(c)
        if groups and groups[-1][0] == chinese:
            groups[-1] = (chinese, groups[-1][1] + c)
        else:
            groups.append((chinese, c))
    return groups


@dataclass
class Element:
    kind: str
    attr: dict[str, Any]
    name: str


class Scene:
    def __init__(self, width=W, height=H, font_dir=Path('/tmp/fonts')):
        self.width, self.height = width, height
        self.font_dir = Path(font_dir)
        self.elements: list[Element] = []
        self.fonts = {}
        self._load_fonts()

    def _load_fonts(self):
        for bold, wt in [(False, '400'), (True, '700')]:
            zh = list(self.font_dir.glob(f'*chinese-simplified-{wt}.ttf'))
            en = self.font_dir / ('DejaVuSans-Bold.ttf' if bold else 'DejaVuSans.ttf')
            if not zh or not en.exists():
                raise FileNotFoundError('Point --font-dir at Noto Sans SC (400/700) and DejaVu Sans TTF files.')
            self.fonts[(True, bold)] = str(zh[0])
            self.fonts[(False, bold)] = str(en)

    def measure(self, s, size, bold=False):
        total = 0.0
        for chinese, part in font_runs(s):
            font = ImageFont.truetype(self.fonts[(chinese, bold)], int(round(size * 4)))
            total += font.getlength(part) / 4
        return total

    def _add(self, kind, name='', **attrs):
        if not name:
            name = f'{kind}-{len(self.elements):03}'
        self.elements.append(Element(kind, attrs, name))

    def rect(self, x, y, w, h, fill, stroke=None, sw=1, radius=0, name=''):
        self._add('rect', name, x=x, y=y, w=w, h=h, fill=fill, stroke=stroke, sw=sw, radius=radius)

    def ellipse(self, cx, cy, rx, ry, fill, stroke=None, sw=1, dash=None, name=''):
        self._add('ellipse', name, cx=cx, cy=cy, rx=rx, ry=ry, fill=fill, stroke=stroke, sw=sw, dash=dash)

    def line(self, x1, y1, x2, y2, color, sw=1, dash=None, name=''):
        self._add('line', name, x1=x1, y1=y1, x2=x2, y2=y2, stroke=color, sw=sw, dash=dash)

    def path(self, d, stroke=None, sw=1, fill=None, dash=None, name=''):
        self._add('path', name, d=d, stroke=stroke, sw=sw, fill=fill, dash=dash)

    def text(self, x, y, s, size=24, color=None, bold=False, anchor='start', width=None, name=''):
        color = color or P['ink']
        if width:
            while self.measure(s, size, bold) > width and size > 16:
                size -= 0.5
        self._add('text', name, x=x, y=y, s=s, size=size, color=color,
                  bold=bold, anchor=anchor, tw=self.measure(s, size, bold))

    def arrowhead(self, x, y, angle, color, size=11, name=''):
        th = math.radians(angle)
        pts = [(x, y)]
        for a in (-0.44, 0.44):
            pts.append((x - size * math.cos(th + a), y - size * math.sin(th + a)))
        d = f'M {pts[0][0]} {pts[0][1]} L {pts[1][0]} {pts[1][1]} L {pts[2][0]} {pts[2][1]} Z'
        self.path(d, fill=color, name=name)

    def arrow(self, d, x, y, angle, color, sw=4, dash=None, name=''):
        self.path(d, stroke=color, sw=sw, dash=dash, name=name)
        self.arrowhead(x, y, angle, color, size=13 if sw >= 4 else 10, name=name + '-箭头')

    def pill(self, cx, cy, w, h, label, fill, stroke, color, size=19, name=''):
        self.rect(cx - w / 2, cy - h / 2, w, h, fill, stroke, 1.5, h / 2, name=name)
        self.text(cx, cy + size * 0.36, label, size, color, True, 'middle', width=w - 12, name=name + '-标签')

    def li(self, x, y, r=20, name='Li+'):
        self.ellipse(x, y, r, r, P['li'], P['li_dark'], 1.5, name=name)
        self.ellipse(x - r * .25, y - r * .30, r * .31, r * .23, '#FFE3A3', name=name + '-高光')
        self.text(x, y + r * .34, 'Li⁺', r * .78, '#FFFFFF', True, 'middle', name=name + '-文字')

    def circle_layers(self, cx, cy, r, palette):
        # Flat concentric fills produce a soft magnifying lens without any
        # bitmap background: editable in PowerPoint, reproducible in SVG.
        for offset, fill in enumerate(palette):
            self.ellipse(cx, cy, r - offset * 3.2, r - offset * 3.2, fill)

    def svg(self, elements=None, view=None):
        elements = elements or self.elements
        if view is None:
            vx, vy, vw, vh = 0, 0, self.width, self.height
        else:
            vx, vy, vw, vh = view
        out = [f'<svg xmlns="http://www.w3.org/2000/svg" width="{vw}" height="{vh}" viewBox="{vx} {vy} {vw} {vh}">',
               '<title>MBA-based GPE design concept — original schematic</title>',
               '<desc>Proposed network, solvation and interphase design. All mechanistic benefits are hypotheses, not measured results. No literature imagery is embedded.</desc>']
        for el in elements:
            a = el.attr
            style = ''
            if el.kind != 'text':
                style += f' fill="{a.get("fill") or "none"}"'
                if a.get('stroke'):
                    style += f' stroke="{a["stroke"]}" stroke-width="{a.get("sw",1)}" stroke-linecap="round" stroke-linejoin="round"'
                if a.get('dash'):
                    style += f' stroke-dasharray="{a["dash"]}"'
            ident = f' data-name="{escape(el.name, quote=True)}"'
            if el.kind == 'rect':
                out.append(f'<rect x="{a["x"]}" y="{a["y"]}" width="{a["w"]}" height="{a["h"]}" rx="{a["radius"]}"{style}{ident}/>')
            elif el.kind == 'ellipse':
                out.append(f'<ellipse cx="{a["cx"]}" cy="{a["cy"]}" rx="{a["rx"]}" ry="{a["ry"]}"{style}{ident}/>')
            elif el.kind == 'line':
                out.append(f'<line x1="{a["x1"]}" y1="{a["y1"]}" x2="{a["x2"]}" y2="{a["y2"]}"{style}{ident}/>')
            elif el.kind == 'path':
                out.append(f'<path d="{a["d"]}"{style}{ident}/>')
            elif el.kind == 'text':
                pieces = []
                for chinese, part in font_runs(a['s']):
                    family = 'Noto Sans SC' if chinese else 'DejaVu Sans'
                    pieces.append(f'<tspan font-family="{family}">{escape(part)}</tspan>')
                out.append(f'<text x="{a["x"]}" y="{a["y"]}" font-size="{a["size"]}" font-weight="{700 if a["bold"] else 400}" fill="{a["color"]}" text-anchor="{a["anchor"]}"{ident}>{"".join(pieces)}</text>')
        out.append('</svg>')
        return '\n'.join(out)

    def validate(self):
        for e in self.elements:
            a = e.attr
            if e.kind == 'text':
                left = a['x'] - (a['tw'] if a['anchor'] == 'end' else a['tw'] / 2 if a['anchor'] == 'middle' else 0)
                assert left >= -1, (e.name, left)
                assert left + a['tw'] < self.width + 1, (e.name, left + a['tw'])
                assert a['y'] - a['size'] >= 0 and a['y'] <= self.height, e.name
            elif e.kind == 'rect':
                assert a['w'] > 0 and a['h'] > 0
                assert a['x'] >= 0 and a['y'] >= 0
                assert a['x'] + a['w'] <= self.width + 1 and a['y'] + a['h'] <= self.height + 1, e.name

    def powerpoint(self, path: Path, notes: str):
        prs = Presentation()
        prs.slide_width = Inches(13.3333333333)
        prs.slide_height = Inches(7.5)
        prs.core_properties.title = '体系设计思路：网络—溶剂化—界面协同调控'
        prs.core_properties.subject = 'MBA / LiPF6–LiDFOB / DMTFA:TTE — proposed GPE design'
        prs.core_properties.author = 'Original schematic for the research project'
        prs.core_properties.keywords = 'MBA, GPE, DMTFA, TTE, LiDFOB, design hypothesis, editable vector'
        sl = prs.slides.add_slide(prs.slide_layouts[6])
        k = prs.slide_width / self.width
        emu = lambda v: int(round(v * k))

        def style(sh, a):
            if hasattr(sh, 'fill'):
                if a.get('fill'):
                    sh.fill.solid()
                    sh.fill.fore_color.rgb = rgb(a['fill'])
                else:
                    sh.fill.background()
            if a.get('stroke'):
                sh.line.color.rgb = rgb(a['stroke'])
                sh.line.width = Pt(a.get('sw', 1) / 2)
                sh.line._get_or_add_ln().set('cap', 'rnd')
                if a.get('dash'):
                    sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            else:
                sh.line.fill.background()

        for el in self.elements:
            a = el.attr
            if el.kind == 'rect':
                kind = MSO_SHAPE.ROUNDED_RECTANGLE if a['radius'] else MSO_SHAPE.RECTANGLE
                sh = sl.shapes.add_shape(kind, emu(a['x']), emu(a['y']), emu(a['w']), emu(a['h']))
                if a['radius']:
                    sh.adjustments[0] = min(0.5, a['radius'] / min(a['w'], a['h']))
                style(sh, a)
            elif el.kind == 'ellipse':
                sh = sl.shapes.add_shape(MSO_SHAPE.OVAL, emu(a['cx']-a['rx']), emu(a['cy']-a['ry']), emu(a['rx']*2), emu(a['ry']*2))
                style(sh, a)
            elif el.kind == 'line':
                # Native connectors avoid zero-extent custom geometries for
                # vertical/horizontal lines and remain easily editable.
                sh = sl.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                    emu(a['x1']), emu(a['y1']), emu(a['x2']), emu(a['y2']))
                style(sh, a)
            elif el.kind == 'path':
                paths = flatten_svg_path(a['d'])
                pts, close = paths[0]
                ff = sl.shapes.build_freeform(pts[0][0], pts[0][1], scale=k)
                ff.add_line_segments(pts[1:], close=close)
                for pts, close in paths[1:]:
                    ff.move_to(pts[0][0], pts[0][1])
                    ff.add_line_segments(pts[1:], close=close)
                sh = ff.convert_to_shape()
                style(sh, a)
            elif el.kind == 'text':
                # The scene stores SVG baseline positions. Align the native
                # text frame on the same baseline with no Office margins.
                tw = max(1, a['tw'])
                x = a['x'] - (tw if a['anchor'] == 'end' else tw/2 if a['anchor'] == 'middle' else 0)
                y = a['y'] - a['size'] * 1.025
                sh = sl.shapes.add_textbox(emu(x - 1), emu(y), emu(tw + 6), emu(a['size'] * 1.37))
                tf = sh.text_frame
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.word_wrap = False
                tf.vertical_anchor = MSO_ANCHOR.TOP
                p = tf.paragraphs[0]
                p.alignment = {'middle': PP_ALIGN.CENTER, 'end': PP_ALIGN.RIGHT}.get(a['anchor'], PP_ALIGN.LEFT)
                p.space_before = p.space_after = Pt(0)
                p.line_spacing = 1.0
                for chinese, part in font_runs(a['s']):
                    r = p.add_run()
                    r.text = part
                    r.font.name = 'Microsoft YaHei' if chinese else 'Arial'
                    r.font.bold = a['bold']
                    r.font.size = Pt(a['size'] / 2)
                    r.font.color.rgb = rgb(a['color'])
                    if chinese:
                        rpr = r._r.get_or_add_rPr()
                        ea = OxmlElement('a:ea')
                        ea.set('typeface', 'Microsoft YaHei')
                        rpr.append(ea)
            else:
                raise ValueError(el.kind)
            sh.name = el.name
        sl.notes_slide.notes_text_frame.text = notes
        prs.save(path)
        return prs


def flatten_svg_path(d, segments=28):
    """Convert the small absolute M/L/C/Q/Z vocabulary to editable PPT paths."""
    toks = re.findall(r'[MLCQZmlcqz]|[-+]?(?:\d*\.\d+|\d+)(?:[eE][-+]?\d+)?', d)
    i, cmd = 0, None
    cur = (0., 0.)
    sub = []
    paths = []
    while i < len(toks):
        if toks[i].isalpha():
            cmd = toks[i].upper()
            i += 1
        if cmd == 'Z':
            paths.append((sub, True))
            sub = []
            cmd = None
            continue
        n = {'M': 2, 'L': 2, 'C': 6, 'Q': 4}[cmd]
        vals = [float(t) for t in toks[i:i+n]]
        i += n
        if cmd == 'M':
            if sub:
                paths.append((sub, False))
            cur = tuple(vals)
            sub = [cur]
            cmd = 'L'
        elif cmd == 'L':
            cur = tuple(vals)
            sub.append(cur)
        elif cmd == 'C':
            p0 = cur
            p1, p2, p3 = (vals[:2], vals[2:4], vals[4:6])
            for j in range(1, segments + 1):
                t = j / segments
                u = 1-t
                sub.append(tuple(u**3*p0[z] + 3*u*u*t*p1[z] + 3*u*t*t*p2[z] + t**3*p3[z] for z in (0,1)))
            cur = tuple(p3)
        elif cmd == 'Q':
            p0 = cur
            p1, p2 = (vals[:2], vals[2:4])
            for j in range(1, segments + 1):
                t = j / segments
                u = 1-t
                sub.append(tuple(u*u*p0[z] + 2*u*t*p1[z] + t*t*p2[z] for z in (0,1)))
            cur = tuple(p2)
    if sub:
        paths.append((sub, False))
    return paths


def draw_electrode(s, x, y, h, side, stable=False, label='正极', prefix=''):
    w = 28
    color = P['electrode_c'] if label == '正极' else P['electrode_a']
    s.rect(x, y, w, h, color, radius=3, name=prefix+'电极主体')
    s.rect(x+2, y+2, 4, h-4, '#8999AB' if label=='正极' else '#D5C1AC', radius=2, name=prefix+'电极高光')
    filmx = x+w+5 if side=='left' else x-29
    if stable:
        # LiF/B–O mosaic is a visual design goal, not a compositional map.
        s.rect(filmx, y, 24, h, '#DAEEE8', P['mint'], 1, name=prefix+'连续界面膜底层')
        colors = [P['gold'], P['mint'], '#D8DBE7', P['mint'], P['gold'], '#B0D5C9']
        for j in range(14):
            for col in range(2):
                s.rect(filmx+col*12+1, y+j*(h/14)+1, 10.5, h/14-2,
                       colors[(j+col*2) % len(colors)], name=prefix+f'成膜组分-{j}-{col}')
        s.line(filmx - 2 if side=='right' else filmx+26, y,
               filmx - 2 if side=='right' else filmx+26, y+h, P['teal'], 2.5,
               name=prefix+'均匀界面边界')
    else:
        # Film segments and gaps, not borrowed microscopy or measured data.
        s.rect(filmx+8, y, 9, h, '#D3E3F1', name=prefix+'常规膜底层')
        heights = [28, 17, 37, 21, 32, 16, 25, 40, 21, 31]
        yy = y
        for j, hh in enumerate(heights):
            if yy+hh > y+h:
                hh = y+h-yy
            if hh <= 0:
                break
            protrusion = [16, 8, 21, 12, 18, 6, 15, 26, 10, 17][j]
            if side=='left':
                pts = f'M {filmx+7} {yy} C {filmx+protrusion+14} {yy+hh*.15} {filmx+protrusion+13} {yy+hh*.82} {filmx+7} {yy+hh}'
                pts += f' L {filmx+4} {yy+hh} L {filmx+4} {yy} Z'
            else:
                pts = f'M {filmx+17} {yy} C {filmx-protrusion+2} {yy+hh*.18} {filmx-protrusion+2} {yy+hh*.8} {filmx+17} {yy+hh}'
                pts += f' L {filmx+21} {yy+hh} L {filmx+21} {yy} Z'
            s.path(pts, fill=P['organic'] if j%3 else '#A7C7E3', name=prefix+f'不均匀膜-{j}')
            yy += hh+4
    s.text(x+w/2, y+h+37, label, 24, P['muted'], True, 'middle', name=prefix+'电极标签')


def draw_old_network(s, cx, cy):
    c = P['poly_old']
    def coord(d):
        # Paths below are already in slide coordinates to keep them editable.
        return d
    s.path(coord('M 328 418 C 370 395 408 421 451 413 C 504 403 552 431 645 411'), c, 6.5, name='传统骨架-上链')
    s.path(coord('M 322 560 C 365 537 415 566 468 546 C 514 530 577 563 648 549'), c, 6.5, name='传统骨架-下链')
    s.path(coord('M 370 390 C 347 444 396 480 371 520 C 351 551 368 586 384 604'), c, 6.5, name='传统骨架-左交联链')
    s.path(coord('M 601 393 C 564 435 608 480 584 532 C 572 558 586 589 603 613'), c, 6.5, name='传统骨架-右交联链')
    for x,y in [(373,418),(374,548),(593,419),(586,549)]:
        s.ellipse(x,y,7.5,7.5,'#FFFFFF',c,2,name='传统交联节点')
    s.ellipse(cx,cy,67,67,'#FFFFFF',P['blue_mid'],1.5,dash='4 5',name='传统溶剂化鞘-示意')
    # Keep the polymer polar site in front of, not hidden under, the shell.
    s.path('M 376 476 C 389 471 399 470 409 473', c, 3, name='传统极性基团-支链')
    s.line(409,470,420,474,P['coral'],2.5,name='传统羰基-双键1')
    s.line(408,476,419,480,P['coral'],2.5,name='传统羰基-双键2')
    s.ellipse(427,479,9,9,P['coral'],name='传统羰基氧')
    s.text(427,484,'O',14,'#FFFFFF',True,'middle',name='传统氧-标记')
    s.text(395,450,'C=O',18,P['coral'],True,'middle',name='传统极性基团-说明')
    s.line(438,480,cx-21,cy,P['coral'],2.8,dash='4 4',name='传统Li配位束缚-示意')
    solvent_pos=[(cx-15,cy-44),(cx+46,cy-19),(cx+24,cy+44),(cx-40,cy+28)]
    for j,(x,y) in enumerate(solvent_pos):
        s.ellipse(x,y,20,13,P['blue_mid'],'#A8C5E1',1.2,name=f'常规溶剂-{j}')
        s.text(x,y+5,'S',15,P['blue'],True,'middle',name=f'常规溶剂-标签-{j}')
        s.line(cx+(x-cx)*.37,cy+(y-cy)*.37,x-(x-cx)*.30,y-(y-cy)*.30,P['blue'],1.5,dash='4 4',name=f'常规溶剂配位-{j}')
    s.li(cx,cy,20,'传统Li离子')
    s.text(cx,cy+91,'强束缚 / 交换受限',23,P['blue'],True,'middle',name='传统体相瓶颈')


def draw_mba_network(s, cx, cy):
    t = P['teal']
    s.path('M 1323 390 C 1310 428 1338 462 1322 502 C 1308 544 1323 579 1318 610',t,6.5,name='MBA网络-左聚合物主链')
    s.path('M 1551 391 C 1567 427 1538 465 1557 503 C 1573 543 1551 580 1558 608',t,6.5,name='MBA网络-右聚合物主链')
    s.path('M 1319 573 C 1360 561 1401 582 1440 568 C 1485 554 1530 580 1559 569',t,5.5,name='MBA网络-下交联连接')
    # Correct, explicitly labelled MBA amide bridge after polymerization:
    # backbone–C(=O)–NH–CH2–NH–C(=O)–backbone.
    # Carbonyl double bonds terminate on C, never on N.
    s.line(1323,411,1339,411,t,3,name='MBA酰胺桥-左端')
    s.line(1535,411,1554,411,t,3,name='MBA酰胺桥-右端')
    for x, label, nm in [(1349,'C','左羰基碳'),(1387,'NH','左酰胺'),
                          (1437,'CH₂','亚甲基'),(1487,'NH','右酰胺'),(1525,'C','右羰基碳')]:
        s.text(x,419,label,21,P['teal_dark'],True,'middle',name='MBA酰胺桥-'+nm)
    for xa, xb in [(1358,1371),(1403,1418),(1456,1471),(1503,1516)]:
        s.line(xa,411,xb,411,t,2,name='MBA酰胺桥-单键')
    for x,nm in [(1349,'左'),(1525,'右')]:
        s.line(x-3,398,x-3,384,P['coral'],2,name='MBA羰基-'+nm+'双键1')
        s.line(x+3,398,x+3,384,P['coral'],2,name='MBA羰基-'+nm+'双键2')
        s.text(x,379,'O',21,P['coral'],True,'middle',name='MBA羰基氧-'+nm)
    for x,y in [(1323,411),(1553,411),(1319,573),(1559,569)]:
        s.ellipse(x,y,7.5,7.5,'#FFFFFF',t,2,name='MBA交联节点')
    s.ellipse(cx,cy,79,79,'#FFFFFF','#EAD0C6',1.5,dash='5 5',name='候选溶剂化鞘-非实测配位数')
    # Species are pictograms, not an atomistic MD snapshot. Connecting dashes
    # denote proposed coordination/exchange, not a claimed coordination number.
    ligands=[(1392,465,'DMTFA',84,30,P['solvent'],P['solvent_line'],P['coral']),
             (1468,542,'DMTFA',84,30,P['solvent'],P['solvent_line'],P['coral']),
             (1497,470,'DFOB⁻',76,31,P['dfob'],'#4C947F','#FFFFFF'),
             (1389,539,'PF₆⁻',64,30,P['pf'],'#8F85B7','#FFFFFF')]
    for j,(x,y,l,w,h,f,st,col) in enumerate(ligands):
        dx,dy=x-cx,y-cy
        s.line(cx+dx*.35,cy+dy*.35,x-dx*.23,y-dy*.23,P['coral'] if l=='DMTFA' else P['teal'],1.8,dash='4 4',name=f'本体系候选配位-{l}-{j}')
        s.pill(x,y,w,h,l,f,st,col,17,name=f'本体系配位物种-{l}-{j}')
    s.li(cx,cy,20,'本体系Li离子')
    # Illustrative hydrogen bond from MBA N–H to the DMTFA carbonyl oxygen;
    # existence/strength and the resulting miscibility remain to be verified.
    s.line(1392,427,1392,447,P['coral'],1.8,dash='3 4',name='MBA-DMTFA候选氢键')
    s.ellipse(1392,448,4,4,P['coral'],name='DMTFA候选氢键受体-氧位点')
    s.text(1359,443,'氢键',16,P['coral'],False,'middle',name='MBA-DMTFA氢键-说明')
    # TTE is drawn outside the illustrative first shell: a low-coordination
    # diluent role, not an absolute claim of zero TTE coordination.
    for j,(x,y) in enumerate([(1328,475),(1563,525),(1437,628)]):
        s.pill(x,y,65,28,'TTE',P['tte'],P['tte_line'],'#527F83',17,name=f'TTE稀释剂-外侧-{j}')
    s.text(cx,cy+98,'溶剂 / 阴离子协同配位',22,P['teal_dark'],True,'middle',name='候选溶剂化结构-说明')
    s.arrow('M 1370 498 C 1377 516 1400 525 1419 516',1419,516,-20,P['teal'],2.5,name='动态配位交换-示意')


def build(font_dir):
    s = Scene(font_dir=font_dir)
    s.rect(0,0,W,H,P['bg'],name='页面背景')
    s.rect(50,36,7,65,P['teal'],radius=3,name='标题强调线')
    s.text(76,59,'DESIGN CONCEPT',20,P['teal_dark'],True,name='英文页眉')
    s.text(76,111,'体系设计思路：网络—溶剂化—界面协同调控',49,P['ink'],True,width=1785,name='主标题')
    s.rect(50,142,1820,62,P['white'],P['line'],1.2,14,name='配方栏')
    s.text(74,183,'MBA 交联网络',29,P['teal_dark'],True,name='配方-MBA')
    s.line(317,158,317,187,P['line'],1.5,name='配方分隔线-1')
    s.text(344,183,'1.0 M LiPF₆ + 0.2 M LiDFOB',29,P['ink'],True,name='配方-双盐浓度')
    s.line(828,158,828,187,P['line'],1.5,name='配方分隔线-2')
    s.text(855,183,'DMTFA : TTE = 1 : 1',29,P['ink'],True,name='配方-溶剂比例')
    s.pill(1723,173,246,34,'设计预期 · 待验证',P['coral_light'],'#E9C9BE',P['coral'],21,name='设计假设标记')

    panel_start = len(s.elements)
    s.rect(50,230,875,511,P['white'],P['line'],1.5,20,name='传统体系面板')
    s.rect(995,230,875,511,P['white'],P['line'],1.5,20,name='本体系设计面板')
    s.rect(50,230,875,74,P['blue_light'],radius=20,name='传统面板标题底色')
    s.rect(50,267,875,37,P['blue_light'],name='传统面板标题底色-平底')
    s.rect(995,230,875,74,P['coral_light'],radius=20,name='设计面板标题底色')
    s.rect(995,267,875,37,P['coral_light'],name='设计面板标题底色-平底')
    s.text(79,278,'传统凝胶',35,P['blue'],True,name='传统面板标题')
    s.text(895,277,'典型瓶颈（示意）',24,P['muted'],False,'end',name='传统面板提示')
    s.text(1024,278,'本工作：MBA 协同凝胶体系',34,P['teal_dark'],True,name='本体系面板标题')
    s.text(1840,277,'设计预期',23,P['coral'],True,'end',name='本体系面板提示')

    # Main comparison: no chemistry or electrode identity is copied from the
    # source figure (its HV-LCO, Li metal, E/F-GPE, TFSI, FDMA do not appear).
    draw_electrode(s,85,360,273,'left',False,'正极','传统-正极-')
    draw_electrode(s,865,360,273,'right',False,'负极','传统-负极-')
    draw_electrode(s,1030,360,273,'left',True,'正极','设计-正极-')
    draw_electrode(s,1810,360,273,'right',True,'负极','设计-负极-')

    s.text(188,340,'不均匀 CEI',22,P['blue'],True,'middle',name='传统正极界面标签')
    s.text(793,340,'不均匀 SEI',22,P['blue'],True,'middle',name='传统负极界面标签')
    s.text(1150,340,'LiF / B–O CEI',22,P['teal_dark'],True,'middle',name='预期CEI组分')
    s.text(1739,340,'LiF / B–O SEI',22,P['teal_dark'],True,'middle',name='预期SEI组分')

    s.circle_layers(487,486,185,['#E7EFF7','#EAF1F8','#EDF3F9','#F1F6FA','#F5F8FB','#F8FAFC'])
    s.circle_layers(1437,486,185,['#F9E4DE','#F9E9E3','#FBECE7','#FCF0EB','#FDF4EF','#FEF7F4'])
    s.text(487,366,'常规聚合物网络',23,P['blue'],True,'middle',name='传统网络说明')
    s.text(1437,360,'MBA 酰胺交联网络',23,P['teal_dark'],True,'middle',name='MBA网络说明')
    draw_old_network(s,487,486)
    draw_mba_network(s,1437,499)

    # Qualitative kinetic arrows are not quantitative energy-barrier plots.
    s.arrow('M 157 517 C 173 517 180 517 185 499 C 199 452 210 448 223 494 C 233 530 248 532 273 512',273,512,-28,P['blue'],4.5,'7 6',name='传统-迁移受限箭头')
    s.text(216,432,'迁移受限',22,P['blue'],True,'middle',name='传统-传输标注')
    s.arrow('M 684 518 C 706 519 715 518 722 490 C 735 442 749 443 759 493 C 767 532 797 532 827 511',827,511,-20,P['blue'],4.5,'7 6',name='传统-脱溶剂化受阻箭头')
    s.text(754,430,'脱溶剂化受阻',21,P['blue'],True,'middle',width=156,name='传统-脱溶剂化标注')
    s.arrow('M 1095 512 C 1122 512 1137 503 1160 501 C 1180 499 1193 510 1220 508',1220,508,0,P['teal'],4.5,'7 6',name='设计-配位交换箭头')
    s.text(1160,432,'促进动态交换',21,P['teal_dark'],True,'middle',width=153,name='设计-交换标注')
    s.arrow('M 1651 512 C 1673 512 1684 500 1704 500 C 1730 500 1748 512 1774 509',1774,509,0,P['coral'],4.5,'7 6',name='设计-脱溶剂化调控箭头')
    s.text(1704,432,'脱溶剂化调控',21,P['coral'],True,'middle',width=148,name='设计-脱溶剂化标注')

    s.text(211,605,'界面副反应',21,P['muted'],False,'middle',name='传统正极-副反应')
    s.text(754,605,'局部通量不均',21,P['muted'],False,'middle',name='传统负极-通量')
    s.text(1151,605,'界面成膜调控',21,P['teal_dark'],False,'middle',name='设计正极-成膜')
    s.text(1704,605,'均匀 Li⁺ 通量',20,P['teal_dark'],False,'middle',width=148,name='设计负极-通量')
    s.line(77,687,898,687,P['line'],1.2,name='传统总结分隔线')
    s.line(1022,687,1843,687,P['line'],1.2,name='设计总结分隔线')
    s.text(487,721,'配位束缚 + 非均匀传输 + 界面失稳',26,P['blue'],True,'middle',name='传统瓶颈总结')
    s.text(1437,721,'锁液稳网 + 动态配位 + 双界面稳定',26,P['teal_dark'],True,'middle',name='设计目标总结')
    # Center transition icon.
    s.ellipse(960,484,24,24,'#E5EEED',name='对比转向圆底')
    s.line(947,484,970,484,P['teal'],3,name='对比转向箭杆')
    s.arrowhead(974,484,0,P['teal'],12,name='对比转向箭头')

    # Compact legend, six groups, all editable shapes.
    y=784
    s.line(76,y,114,y,P['teal'],5,name='图例-MBA网络')
    s.text(127,y+7,'MBA 网络',23,P['muted'],name='图例-MBA文字')
    s.li(302,y,12,'图例-Li')
    s.text(323,y+7,'Li⁺',23,P['muted'],name='图例-Li文字')
    s.ellipse(453,y,25,11,P['solvent'],P['solvent_line'],1.2,name='图例-DMTFA')
    s.text(490,y+7,'DMTFA',23,P['muted'],name='图例-DMTFA文字')
    s.rect(677,y-11,45,22,P['tte'],P['tte_line'],1.2,11,name='图例-TTE')
    s.text(736,y+7,'TTE',23,P['muted'],name='图例-TTE文字')
    s.ellipse(903,y,12,12,P['pf'],name='图例-PF6')
    s.ellipse(934,y,12,12,P['dfob'],name='图例-DFOB')
    s.text(958,y+7,'PF₆⁻ / DFOB⁻',23,P['muted'],name='图例-阴离子文字')
    s.rect(1267,y-11,22,22,P['gold'],radius=3,name='图例-LiF')
    s.rect(1296,y-11,22,22,P['mint'],radius=3,name='图例-BO')
    s.text(1331,y+7,'LiF / B–O 成膜组分（预期）',23,P['muted'],name='图例-界面组分文字')
    panel_end = len(s.elements)

    # Three design levers: one chemistry/fact row plus one concise hypothesis.
    card_w=(1820-26*2)/3
    card_x=[50+j*(card_w+26) for j in range(3)]
    headings=['MBA：交联锁液与相容','DMTFA/TTE：溶剂化调控','LiDFOB：双界面成膜调控']
    second=['CH₂=CH–CO–NH–CH₂–NH–CO–CH=CH₂',
            'DMTFA (CF₃CON(CH₃)₂) + TTE',
            '阴离子参与溶剂化与界面反应']
    third=['双键交联；N–H···O=C 辅助相容',
           '调节配位平衡，促进 Li⁺ 动态交换',
           '预期构筑 LiF / B–O 富集 SEI / CEI']
    colors=[P['teal'],P['coral'],P['blue']]
    for j,x in enumerate(card_x):
        s.rect(x,824,card_w,171,P['white'],P['line'],1.5,16,name=f'设计抓手-{j+1}-卡片')
        s.rect(x,842,5,132,colors[j],radius=2,name=f'设计抓手-{j+1}-强调线')
        s.pill(x+47,857,42,30,f'0{j+1}',P['teal_light'] if j==0 else P['coral_light'] if j==1 else P['blue_light'],
               P['teal_light'] if j==0 else P['coral_light'] if j==1 else P['blue_light'],colors[j],19,name=f'设计抓手-{j+1}-编号')
        s.text(x+83,868,headings[j],29,colors[j],True,width=card_w-106,name=f'设计抓手-{j+1}-标题')
        s.text(x+25,918,second[j],24,P['muted'],False,width=card_w-45,name=f'设计抓手-{j+1}-化学依据')
        s.text(x+25,966,third[j],25,P['ink'],True,width=card_w-45,name=f'设计抓手-{j+1}-作用目标')

    s.text(50,1043,'注：图示为设计机理假设，非实测配位数、能垒或界面组成；需结合光谱/模拟、传输测试及 XPS 等验证。',21,P['muted'],False,width=1800,name='科学边界说明')
    s.validate()
    ET.fromstring(s.svg())
    return s, panel_start, panel_end


NOTES = '''【体系设计页：建议紧接“传统凝胶电解质共性问题”页】

配方：MBA（N,N′-亚甲基双丙烯酰胺）作为交联单体；
液态组成为 1.0 M LiPF6 + 0.2 M LiDFOB in DMTFA:TTE = 1:1。
用户未说明溶剂比的体积/质量/摩尔基准、MBA含量、引发剂及成胶条件，图中不作补充假定。

【约1–2分钟讲稿】
上一页提出了凝胶化后传输受限与电极界面失稳两类问题。针对这些瓶颈，我们从网络、溶剂化和成膜三个层次进行协同设计。
首先，MBA的两个可聚合双键用于构建三维交联网络，通过适当的交联程度锁住电解液；酰胺基团的分子间作用则有望改善网络与混合溶剂的相容性。我们需要平衡保液能力与离子传输，避免交联过度。
其次，选择DMTFA/TTE混合体系，调节溶剂、聚合物极性位点和阴离子对Li+的配位平衡，探索动态配位交换与脱溶剂化的改善，而不是简单把Li+固定在骨架上。
最后，在LiPF6基础上引入LiDFOB，利用含氟/含硼组分的界面反应，尝试构建富LiF及B–O组分的SEI/CEI，降低持续副反应。
因此，这页右图表达的是我们的设计目标：锁液稳网、动态配位与双界面稳定。具体机制和性能提升将在后续通过实验和计算验证。

【科学边界与验证】
1. 所有右图的“动态交换”“均匀通量”“LiF/B–O富集膜”均是设计预期，不是当前体系已经得到证明的结果。
2. 不因DMTFA:TTE=1:1或总盐浓度1.2M就直接认定局部高浓（LHCE）、CIP/AGG占优或阴离子配位数增加。
3. 弱溶剂配位不必然带来更低的总脱溶剂化能垒；Li–阴离子和Li–聚合物配位可能造成额外限制。
4. C=O是候选Li+配位位点；N–H可参与氢键，但“锚定阴离子”“提高tLi+”不能仅凭官能团存在作结论。
5. 不宣称LiDFOB清除HF、具体优先氧化/还原次序或确切LUMO/HOMO排序，除非已有对应证据。
6. 正/负极未指定材料，因此使用通用电极示意；没有照搬参考图的HV-LCO、锂金属、FDMA、TFSI或E/F-GPE命名。
7. 网络部分的酰胺桥表示MBA聚合后的骨架侧链连接，并非完整聚合物重复单元/实际网络拓扑。

建议验证：
网络/相容：FTIR或NMR、凝胶分数、保液/溶胀、流变或力学；与交联度的关系。
溶剂化/传输：Raman/FTIR/NMR与MD/DFT、温度依赖电导率、迁移数；EIS及合适的动力学分析。
双界面：XPS/深度分析或TOF-SIMS、循环前后EIS、形貌与全电池/对称电池测试，按实际电极体系选择。

编辑说明：
该PPT为单页16:9原生形状版；所有标题、分子符号、网络、箭头、电极与膜组分均可在PowerPoint中选择修改。未嵌入任何文献照片或生成的SEM图。PNG为高清预览；SVG为可编辑矢量源。
'''


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--font-dir',type=Path,default=Path('/tmp/fonts'))
    args=parser.parse_args()
    s,start,end=build(args.font_dir)
    base=OUT/'MBA_GPE_Design_Concept_Slide'
    svg=s.svg()
    base.with_suffix('.svg').write_text(svg,encoding='utf-8')
    base.with_suffix('.png').write_bytes(resvg_py.svg_to_bytes(svg_string=svg,width=3840,height=2160,font_dirs=[str(args.font_dir)]))
    prs=s.powerpoint(base.with_suffix('.pptx'),NOTES)
    figsvg=s.svg(s.elements[start:end],view=(50,230,1820,580))
    figbase=OUT/'MBA_GPE_Design_Comparison_Figure'
    figbase.with_suffix('.svg').write_text(figsvg,encoding='utf-8')
    figbase.with_suffix('.png').write_bytes(resvg_py.svg_to_bytes(svg_string=figsvg,width=3640,height=1160,font_dirs=[str(args.font_dir)]))
    assert len(prs.slides)==1
    assert not any(sh.shape_type == 13 for sh in prs.slides[0].shapes), 'PPT must remain natively editable, no picture layers.'
    print(f'Created original slide: {len(s.elements)} native elements, 3840×2160 PNG, SVG and 1-slide PPTX.')
    print(f'Shape count: {len(prs.slides[0].shapes)}. Comparison-only crop: 3640×1160.')
    print(base.with_suffix('.pptx'))


if __name__=='__main__':
    main()
