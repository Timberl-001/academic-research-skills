#!/usr/bin/env python3
"""Structural checks for the delivered editable GPE design slide.

Run without the rendering fonts:
    python3 -m unittest discover -s gpe_scientific_questions -p 'test_design_concept.py' -v

These checks validate the saved files, not the chemistry or PowerPoint's visual
rendering. The proposed mechanisms still require experimental verification.
"""
from pathlib import Path
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


HERE = Path(__file__).resolve().parent
BASE = HERE / "MBA_GPE_Design_Concept_Slide"
SVG_NS = {"s": "http://www.w3.org/2000/svg"}


class DesignConceptSlideTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.prs = Presentation(BASE.with_suffix(".pptx"))
        cls.slide = cls.prs.slides[0]
        cls.text = "\n".join(
            shape.text for shape in cls.slide.shapes if shape.has_text_frame
        )

    def test_single_widescreen_slide(self):
        self.assertEqual(len(self.prs.slides), 1)
        self.assertAlmostEqual(
            self.prs.slide_width / self.prs.slide_height, 16 / 9, places=6
        )

    def test_native_editable_objects_not_a_picture(self):
        types = {shape.shape_type for shape in self.slide.shapes}
        self.assertTrue(
            types <= {
                MSO_SHAPE_TYPE.AUTO_SHAPE,
                MSO_SHAPE_TYPE.TEXT_BOX,
                MSO_SHAPE_TYPE.FREEFORM,
                MSO_SHAPE_TYPE.LINE,
            },
            f"Unexpected object types: {types}",
        )
        self.assertGreater(len(self.slide.shapes), 200)
        self.assertGreater(
            sum(shape.has_text_frame and bool(shape.text) for shape in self.slide.shapes),
            50,
        )
        with ZipFile(BASE.with_suffix(".pptx")) as archive:
            self.assertIsNone(archive.testzip())
            self.assertFalse(any(n.startswith("ppt/media/") for n in archive.namelist()))

    def test_all_objects_stay_on_slide(self):
        for shape in self.slide.shapes:
            with self.subTest(shape=shape.name):
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, self.prs.slide_width)
                self.assertLessEqual(shape.top + shape.height, self.prs.slide_height)

    def test_user_formulation_and_design_hypothesis_are_present(self):
        for text in (
            "MBA 交联网络",
            "1.0 M LiPF₆ + 0.2 M LiDFOB",
            "DMTFA : TTE = 1 : 1",
            "设计预期",
            "非实测配位数",
            "LiF / B–O",
        ):
            with self.subTest(text=text):
                self.assertIn(text, self.text)
        for copied_label in ("HV-LCO", "FDMA", "TFSI", "E-GPE", "F-GPE"):
            self.assertNotIn(copied_label, self.text)

    def test_speaker_notes_include_scientific_boundaries(self):
        notes = self.slide.notes_slide.notes_text_frame.text
        self.assertIn("讲稿", notes)
        self.assertIn("设计预期", notes)
        self.assertIn("局部高浓", notes)
        self.assertIn("不因", notes)

    def test_svg_and_ppt_share_the_scene(self):
        root = ET.parse(BASE.with_suffix(".svg")).getroot()
        svg_names = [e.attrib["data-name"] for e in root if "data-name" in e.attrib]
        self.assertEqual(svg_names, [shape.name for shape in self.slide.shapes])
        self.assertFalse(root.findall(".//s:image", SVG_NS))
        self.assertEqual(root.attrib["viewBox"], "0 0 1920 1080")

    def test_high_resolution_previews(self):
        for name, size in (
            ("MBA_GPE_Design_Concept_Slide.png", (3840, 2160)),
            ("MBA_GPE_Design_Comparison_Figure.png", (3640, 1160)),
        ):
            with self.subTest(name=name):
                with Image.open(HERE / name) as image:
                    self.assertEqual(image.size, size)
                    image.verify()


if __name__ == "__main__":
    unittest.main()
