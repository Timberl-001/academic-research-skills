#!/usr/bin/env python3
"""Checks for the natively editable traditional-GPE page and download routes."""
from pathlib import Path
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from serve_ppt_download import PAGE, STEM, response_for

HERE = Path(__file__).resolve().parent
BASE = HERE / STEM
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


class TraditionalEditableTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.prs = Presentation(BASE.with_suffix(".pptx"))
        cls.slide = cls.prs.slides[0]

    def test_single_widescreen_page_with_independent_objects(self):
        self.assertEqual(len(self.prs.slides), 1)
        self.assertAlmostEqual(self.prs.slide_width / self.prs.slide_height, 16 / 9, places=6)
        self.assertEqual(len(self.slide.shapes), 206)
        self.assertGreater(sum(s.has_text_frame and bool(s.text) for s in self.slide.shapes), 50)

    def test_native_objects_not_images_or_locked_shapes(self):
        allowed = {MSO_SHAPE_TYPE.AUTO_SHAPE, MSO_SHAPE_TYPE.TEXT_BOX,
                   MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.LINE}
        self.assertTrue(all(s.shape_type in allowed for s in self.slide.shapes))
        with ZipFile(BASE.with_suffix(".pptx")) as archive:
            self.assertIsNone(archive.testzip())
            self.assertFalse(any(n.startswith("ppt/media/") for n in archive.namelist()))
            xml = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
            self.assertFalse(xml.findall(".//a:blip", NS))
            self.assertFalse(xml.findall(".//a:spLocks", NS))
            self.assertFalse(xml.findall(".//a:cxnSpLocks", NS))

    def test_unique_object_names_and_on_slide_bounds(self):
        names = [s.name for s in self.slide.shapes]
        self.assertEqual(len(names), len(set(names)))
        for shape in self.slide.shapes:
            with self.subTest(shape=shape.name):
                self.assertGreaterEqual(shape.left, 0)
                self.assertGreaterEqual(shape.top, 0)
                self.assertLessEqual(shape.left + shape.width, self.prs.slide_width)
                self.assertLessEqual(shape.top + shape.height, self.prs.slide_height)

    def test_only_common_challenges_with_qualified_claims(self):
        text = "\n".join(s.text for s in self.slide.shapes if s.has_text_frame)
        notes = self.slide.notes_slide.notes_text_frame.text
        meta = self.prs.core_properties
        combined = text + notes + meta.title + meta.subject + meta.keywords
        for forbidden in ("MBA", "DMTFA", "TTE", "LiDFOB", "LHCE"):
            self.assertNotIn(forbidden, combined)
        for required in ("电导率可能下降", "迁移数不必然下降", "界面阻抗增长", "引出体系设计"):
            self.assertIn(required, text)
        self.assertIn("科学表述边界", notes)
        self.assertIn("不表示每一种传统凝胶", notes)

    def test_shared_svg_scene_and_high_resolution_preview(self):
        root = ET.parse(BASE.with_suffix(".svg")).getroot()
        self.assertEqual([e.attrib["data-name"] for e in root if "data-name" in e.attrib],
                         [s.name for s in self.slide.shapes])
        self.assertFalse(any(e.tag.endswith("image") for e in root.iter()))
        self.assertNotIn("MBA", "".join(root.itertext()))
        with Image.open(BASE.with_suffix(".png")) as image:
            self.assertEqual(image.size, (3840, 2160))
            image.verify()

    def test_standalone_bundle_contains_exact_deliverables(self):
        with ZipFile(HERE / f"{STEM}_Files.zip") as archive:
            self.assertIsNone(archive.testzip())
            self.assertEqual(set(archive.namelist()), {
                f"{STEM}.pptx", f"{STEM}.svg", f"{STEM}.png", "README_EDITING.md"
            })
            self.assertEqual(archive.read(f"{STEM}.pptx"), BASE.with_suffix(".pptx").read_bytes())

    def test_download_response_has_attachment_and_correct_mime(self):
        status, payload, headers = response_for(f"/files/{STEM}.pptx?from=button")
        self.assertEqual(status, 200)
        self.assertEqual(payload, BASE.with_suffix(".pptx").read_bytes())
        self.assertEqual(headers["Content-Type"],
                         "application/vnd.openxmlformats-officedocument.presentationml.presentation")
        self.assertEqual(headers["Content-Disposition"], f'attachment; filename="{STEM}.pptx"')
        status, preview, headers = response_for("/preview.png")
        self.assertEqual(status, 200)
        self.assertNotIn("Content-Disposition", headers)
        self.assertTrue(preview.startswith(b"\x89PNG"))

    def test_preview_is_same_origin_and_routes_are_whitelisted(self):
        status, body, headers = response_for("/")
        self.assertEqual(status, 200)
        self.assertIn(f'href="/files/{STEM}.pptx"', body.decode())
        self.assertNotIn("localhost", PAGE)
        self.assertNotIn("127.0.0.1", PAGE)
        self.assertNotIn("github.com", PAGE)
        for unsafe in ("/.git/config", "/../.git/config", "/files/../../README.md",
                       "/files/%2e%2e/README.md", "/build_design_concept.py"):
            self.assertEqual(response_for(unsafe)[0], 404)


if __name__ == "__main__":
    unittest.main()
