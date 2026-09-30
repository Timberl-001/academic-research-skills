#!/usr/bin/env python3
"""Check the saved editable component-role table and its download routes."""
import csv
from pathlib import Path
import unittest
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from build_component_roles_table import HEADERS, ROWS
from serve_ppt_download import COMPONENT_STEM, response_for

HERE = Path(__file__).resolve().parent
BASE = HERE / COMPONENT_STEM
NS = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}


class ComponentRolesTableTest(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.prs = Presentation(BASE.with_suffix(".pptx"))
        cls.slide = cls.prs.slides[0]
        cls.tables = [s.table for s in cls.slide.shapes if s.has_table]

    def test_native_table_not_embedded_picture(self):
        self.assertEqual(len(self.prs.slides), 1)
        self.assertAlmostEqual(self.prs.slide_width/self.prs.slide_height, 16/9, places=6)
        self.assertEqual(len(self.tables), 1)
        self.assertEqual((len(self.tables[0].rows), len(self.tables[0].columns)), (7, 3))
        self.assertFalse(any(s.shape_type == MSO_SHAPE_TYPE.PICTURE for s in self.slide.shapes))
        with ZipFile(BASE.with_suffix(".pptx")) as archive:
            self.assertIsNone(archive.testzip())
            self.assertFalse(any(n.startswith("ppt/media/") for n in archive.namelist()))
            root = ET.fromstring(archive.read("ppt/slides/slide1.xml"))
            self.assertEqual(len(root.findall(".//a:tbl", NS)), 1)
            self.assertFalse(root.findall(".//a:blip", NS))

    def test_all_rows_match_source_content(self):
        table = self.tables[0]
        self.assertEqual(tuple(c.text for c in table.rows[0].cells), HEADERS)
        for row, (label, subtitle, role, why) in enumerate(ROWS, 1):
            self.assertEqual(table.cell(row, 0).text, label+"\n"+subtitle)
            self.assertEqual(table.cell(row, 1).text.replace("\n", ""), role)
            self.assertEqual(table.cell(row, 2).text.replace("\n", ""), why)

    def test_scope_and_fair_comparator_notes(self):
        text = "\n".join(s.text for s in self.slide.shapes if s.has_text_frame)
        notes = self.slide.notes_slide.notes_text_frame.text
        self.assertIn("1.0 M LiPF₆ + 0.2 M LiDFOB", text)
        self.assertIn("DMTFA : TTE = 1 : 1", text)
        self.assertIn("不是本体系同时加入", text)
        self.assertIn("需验证", text)
        self.assertIn("PEGDA分子量", notes)
        self.assertIn("不等于实际交联密度相同", notes)
        self.assertIn("PEGDA性能较差", notes)  # explicitly says this is not implied
        self.assertIn("DMTFA本身无N–H", notes)
        self.assertIn("https://doi.org/10.3390/gels9120975", notes)

    def test_csv_preview_and_table_data_are_consistent(self):
        data = BASE.with_suffix(".csv").read_bytes()
        self.assertTrue(data.startswith(b"\xef\xbb\xbf"))
        with BASE.with_suffix(".csv").open(encoding="utf-8-sig", newline="") as file:
            rows = list(csv.reader(file))
        self.assertEqual(rows[0], list(HEADERS))
        self.assertEqual(len(rows), 7)
        self.assertEqual([r[1:] for r in rows[1:]], [[r[2], r[3]] for r in ROWS])
        root = ET.parse(BASE.with_suffix(".svg")).getroot()
        svg_text = "".join(root.itertext())
        for label, subtitle, role, why in ROWS:
            for text in (label, subtitle, role, why):
                self.assertIn(text, svg_text)
        self.assertFalse(any(e.tag.endswith("image") for e in root.iter()))
        with Image.open(BASE.with_suffix(".png")) as image:
            self.assertEqual(image.size, (3840, 2160))
            image.verify()

    def test_bounds_and_ordered_native_cell_borders(self):
        for shape in self.slide.shapes:
            self.assertGreaterEqual(shape.left, 0)
            self.assertGreaterEqual(shape.top, 0)
            self.assertLessEqual(shape.left + shape.width, self.prs.slide_width)
            self.assertLessEqual(shape.top + shape.height, self.prs.slide_height)
        for row in self.tables[0].rows:
            for cell in row.cells:
                names = [e.tag.split("}")[-1] for e in cell._tc.get_or_add_tcPr()]
                self.assertEqual(names[:4], ["lnL", "lnR", "lnT", "lnB"])
                self.assertLessEqual(len(cell.text_frame.paragraphs), 2)

    def test_same_origin_table_page_and_download_headers(self):
        status, body, headers = response_for("/components")
        self.assertEqual(status, 200)
        self.assertIn(f'/files/{COMPONENT_STEM}.pptx', body.decode())
        self.assertNotIn("localhost", body.decode())
        self.assertNotIn("127.0.0.1", body.decode())
        self.assertEqual(response_for("/", default_page="components")[1], body)
        for suffix, mime in ((".pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
                             (".csv", "text/csv; charset=utf-8")):
            status, body, headers = response_for(f"/files/{COMPONENT_STEM}{suffix}")
            self.assertEqual(status, 200)
            self.assertEqual(body, BASE.with_suffix(suffix).read_bytes())
            self.assertEqual(headers["Content-Type"], mime)
            self.assertTrue(headers["Content-Disposition"].startswith("attachment;"))


if __name__ == "__main__":
    unittest.main()
